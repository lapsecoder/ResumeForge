"""Shared test fixtures for ResumeForge ingestion tests.

All fixtures create synthetic content. No real resume data is used.
"""

from __future__ import annotations

import io
import tempfile as _tempfile
import zipfile
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


def install_tempfile_tracker(monkeypatch: Any, created: list[Path]) -> None:
    """Wrap TemporaryDirectory as used by the ingestion service to record dirs.

    Delegates to the real tempfile.TemporaryDirectory (so directories
    genuinely exist and are cleaned up) while capturing each created
    directory path for assertions.
    """

    from app.ingestion import service as service_module

    original = _tempfile.TemporaryDirectory

    def tracking(*args: Any, **kwargs: Any) -> Any:
        td = original(*args, **kwargs)
        created.append(Path(td.name))
        return td

    monkeypatch.setattr(service_module, "TemporaryDirectory", tracking)


def _escape_pdf_text(text: str) -> str:
    """Escape a string for use inside a PDF text-showing operator."""
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _add_pdf_page(writer: PdfWriter, text: str) -> None:
    """Add a page to the writer with the given extractable text."""
    page = writer.add_blank_page(width=612, height=792)

    content = f"BT /F1 12 Tf 72 720 Td ({_escape_pdf_text(text)}) Tj ET".encode(
        "latin-1"
    )
    stream = DecodedStreamObject()
    stream.set_data(content)
    page[NameObject("/Contents")] = stream

    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_ref = writer._add_object(font)
    fonts = DictionaryObject({NameObject("/F1"): font_ref})
    resources = DictionaryObject({NameObject("/Font"): fonts})
    page[NameObject("/Resources")] = resources


def make_pdf_bytes(*pages_text: str) -> bytes:
    """Create a valid multi-page PDF, one page per text string."""
    writer = PdfWriter()
    for text in pages_text:
        _add_pdf_page(writer, text)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def make_empty_pdf() -> bytes:
    """Create a valid PDF with no extractable text (simulates a scanned page)."""
    return make_pdf_bytes("")


def make_docx_bytes(*paragraphs: str) -> bytes:
    """Create a valid DOCX with the given paragraphs."""
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_docx_with_table_bytes() -> bytes:
    """Create a valid DOCX with paragraphs and a table."""
    doc = Document()
    doc.add_paragraph("Header text")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Col A"
    table.cell(0, 1).text = "Col B"
    table.cell(1, 0).text = "Val 1"
    table.cell(1, 1).text = "Val 2"
    doc.add_paragraph("Footer text")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_empty_docx_bytes() -> bytes:
    """Create a valid DOCX with no paragraphs."""
    doc = Document()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_malformed_docx_bytes() -> bytes:
    """Create an invalid DOCX (ZIP with no document.xml)."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>',
        )
    return buf.getvalue()


def make_text_bytes(text: str, encoding: str = "utf-8") -> bytes:
    return text.encode(encoding)


@pytest.fixture
def sample_pdf() -> bytes:
    return make_pdf_bytes("John Doe\nSoftware Engineer")


@pytest.fixture
def sample_pdf_multipage() -> bytes:
    return make_pdf_bytes("Page 1 content", "Page 2 content", "Page 3 content")


@pytest.fixture
def empty_pdf() -> bytes:
    return make_pdf_bytes("")


@pytest.fixture
def sample_docx() -> bytes:
    return make_docx_bytes("John Doe", "Software Engineer with 5 years experience")


@pytest.fixture
def sample_docx_with_table() -> bytes:
    return make_docx_with_table_bytes()


@pytest.fixture
def empty_docx() -> bytes:
    return make_empty_docx_bytes()


@pytest.fixture
def malformed_docx() -> bytes:
    return make_malformed_docx_bytes()


@pytest.fixture
def sample_txt() -> bytes:
    return make_text_bytes("John Doe\nSoftware Engineer")


@pytest.fixture
def sample_txt_crlf() -> bytes:
    return make_text_bytes("John Doe\r\nSoftware Engineer\r\n")


@pytest.fixture
def empty_txt() -> bytes:
    return make_text_bytes("")


@pytest.fixture
def malformed_pdf() -> bytes:
    return make_text_bytes("This is not a PDF file")
